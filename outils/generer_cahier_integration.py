"""CAHIER DE RECETTE DE LA BRANCHE D'INTEGRATION integration/recette-2026-10-29.

Une feuille par module, chacune tenant sur UNE page A4 paysage (et a la largeur d'un ecran 1366 px a 100 %), plus une
feuille Synthese. Les cas portent sur ce que la branche apporte depuis la derniere version de dev (05/10) : ventes,
caisse, differes, assurances, commandes, stock, etiquettes, pilotage, campagnes, RH, support, performances, editions.

    python outils/generer_cahier_integration.py SORTIE.xlsx [RESULTATS.txt]

RESULTATS.txt (facultatif) : sortie de l'executeur des tests d'ecran (lignes « === chemin » puis « --- code N ») ; la
colonne « Test auto » porte alors OK / KO pour les cas couverts. Le PDF se tire du classeur par
outils/exporter_pdf.ps1 (Excel), qui ajuste aussi la hauteur des lignes.
"""
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = sys.argv[1]
RESULTATS = sys.argv[2] if len(sys.argv) > 2 else None
BRANCHE = 'integration/recette-2026-10-29'
POLICE = 'Arial'

# ---------------------------------------------------------------- cas : (module, ecran, cas, etapes, attendu, priorite, tests)
CAS = []
MODULES = OrderedDict()


def module(nom, menu):
    MODULES[nom] = menu
    return nom


def c(m, ecran, cas, etapes, attendu, prio='Moyenne', tests=()):
    CAS.append((m, ecran, cas, etapes, attendu, prio, tuple(tests)))


m = module('Connexion et sécurité', 'Page de connexion · GESTION DES UTILISATEURS')
c(m, 'Connexion', 'Connexion normale', '1. Identifiant, mot de passe, Entrée', 'Accueil ouvert avec les menus du profil', 'Haute')
c(m, 'Connexion', 'Mots de passe faux répétés', '1. 10 essais faux\n2. Attendre, réessayer', 'Refus au-delà de 10 essais, nouvelle série après l\'attente ; message sans dire si le compte existe', 'Haute')
c(m, 'Connexion', 'Connexion par badge QR', '1. « Se connecter avec mon badge »\n2. Scanner le badge', 'Session au nom de l\'employé ; badge régénéré ou désactivé refusé', 'Haute')
c(m, 'Connexion', 'Douchette QWERTY / Verr. Maj.', '1. Scanner le badge avec la douchette mal réglée', 'Connexion acceptée ou message « lecture déformée » explicite', 'Basse')
c(m, 'Connexion', 'Bouton « Pointage »', '1. Cliquer « Pointage » sur la page de connexion', 'Le QR de pointage mobile s\'affiche', 'Basse')
c(m, 'Session', 'Expiration après inactivité', '1. Laisser la session inactive\n2. Agir', 'Retour à la connexion ; aucune page protégée sans session', 'Haute')
c(m, 'Droits', 'Super administrateur invisible', '1. Se connecter en administrateur\n2. Ouvrir la liste des utilisateurs et le Centre de support', 'Super administrateur absent de la liste ; Centre de support (sauf « Me contacter ») et audit réservés', 'Haute')
c(m, 'Droits', 'Droits vérifiés par le serveur', '1. Profil sans le droit (ajustement, réserve, inventaire…)\n2. Tenter l\'action', 'Refus du serveur même si le bouton est appelé directement', 'Haute')
c(m, 'Droits', 'Vente close par un profil prévente', '1. Profil prévente\n2. Tenter de terminer une vente', 'Refus : seul un caissier clôture', 'Haute', ['securite/test-cloture-vente-profil-prevente.js'])
c(m, 'Gestion de profils', 'Créer, modifier, attribuer les privilèges', '1. Créer un profil\n2. Le modifier, attribuer des privilèges', 'Fenêtres modales, enregistrement correct', 'Moyenne')
c(m, 'Utilisateurs', 'Empreinte du mot de passe', '1. Ouvrir la liste des utilisateurs (outils du navigateur)', 'Aucune empreinte de mot de passe envoyée au navigateur', 'Haute')
c(m, 'Menus', 'Libellés et ordre des menus', '1. Parcourir les menus', 'Libellés en MAJUSCULES, 25 caractères au plus, ordre fixe des modules', 'Basse', ['retours/test-menu-ordre-modules.js', 'retours/test-retours-0710b.js'])
c(m, 'Menus', 'Thèmes de menu (Métro, A, C, D)', '1. Changer de thème de menu\n2. Ouvrir des écrans au clavier', 'Thème retenu par utilisateur, menu actuel par défaut, aucun écran disparu', 'Moyenne')
c(m, 'Postes', 'Rappel « date ou heure fausse »', '1. Poste à l\'heure décalée', 'Rappel avec « Mettre à jour » et « Masquer »', 'Basse')

m = module('Ventes', 'SERVICE CLIENT')
c(m, 'Ventes', 'Vente au comptant', '1. Nouvelle vente, scanner, quantité\n2. Terminer, régler en espèces', 'Ticket, stock et caisse à jour', 'Haute', ['ui-caisse-deconditionnement.js'])
c(m, 'Ventes', 'Vente assurance avec remise', '1. Vente assurance, articles, remise\n2. Terminer', 'Remise sur la part client de chaque produit, plafond réparti au prorata', 'Haute', ['audits/repro-vente-assurance.js', 'audits/repro-remise-caisse.js'])
c(m, 'Ventes', 'Quantité décimale', '1. Taper 1.5 en quantité', 'Refus avec message clair, quantité précédente gardée', 'Moyenne', ['journal/test-vente-quantite-decimale.js'])
c(m, 'Ventes', 'Vendeur tapé sans être choisi', '1. Taper un vendeur, Tab\n2. Ajouter un article', 'Vendeur reconnu ou signalé ; le 1er article s\'ajoute', 'Moyenne', ['journal/test-vente-vendeur-saisi.js'])
c(m, 'Ventes', 'QR publicitaire scanné', '1. Scanner un QR de site internet', '« QR code publicitaire » sans fenêtre, aucune recherche', 'Basse', ['journal/test-scan-qr-publicitaire.js'])
c(m, 'Ventes', 'Étiquette GS1 / DataMatrix', '1. Scanner une boîte GS1', 'Produit trouvé, lot contrôlé, péremption effacée après l\'ajout', 'Moyenne', ['retours/test-gs1-bascule.js', 'retours/test-gs1-controle-lot.js'])
c(m, 'Ventes', 'Alerte d\'interactions', '1. Activer l\'alerte\n2. Vendre deux produits qui interagissent', 'Alerte affichée ; coupée par défaut', 'Basse', ['retours/test-interactions-vente.js'])
c(m, 'Ventes', 'Points fidélité', '1. Client fidélité, case « Utiliser »\n2. Régler en points', 'Points déduits, restitués à l\'annulation', 'Moyenne')
c(m, 'Ventes', 'Nouveau client depuis la vente', '1. « Nouveau client » dans la vente\n2. Enregistrer', 'Client créé sans erreur JavaScript', 'Moyenne')
c(m, 'Ventes terminées', 'Modifier une vente', '1. Ventes terminées > Modifier\n2. Changer, terminer', 'Copie modifiée ; vente facturée ou différé réglé : refus avec la raison', 'Haute', ['points-28-10/test-points-29-10.js', 'audits/test-ecran-modifier-vente-deja-modifiee.js'])
c(m, 'Ventes terminées', 'Annuler une vente', '1. Annuler\n2. Annuler de nouveau', 'Stock rendu, une seule annulation ; refus si en cours de modification ou différé réglé', 'Haute', ['audits/repro-vente-comptant.js', 'audits/repro-vo-remise.js'])
c(m, 'Ventes terminées', 'Colonnes et retour', '1. Ouvrir la liste, Modifier puis Retour', 'Date et heure, « Différé », poste ; retour avec les mêmes critères', 'Basse', ['retours/test-ventes-terminees-colonnes.js'])
c(m, 'Prévente', 'Ticket de prévente scanné', '1. Scanner le ticket en caisse', 'Prévente retrouvée (même lu « 261025)00145 ») ; QR dans le rouleau', 'Moyenne')
c(m, 'Vente dépôt', 'Vente à un dépôt', '1. Nouvelle vente dépôt, articles, terminer', 'Stocks officine et dépôt à jour ; historique cohérent', 'Moyenne', ['vente/test-historique-deconditionnement.js', 'retours/test-ca-depot.js'])
c(m, 'Proforma', 'Proforma', '1. Nouveau proforma, client, articles, imprimer', 'Aucun mouvement de stock', 'Basse', ['retours/test-proforma.js'])

m = module('Caisse et gestion courante', 'GESTION COURANTE')
c(m, 'Ouverture / clôture', 'Ouvrir puis clôturer la caisse', '1. Ouvrir avec le fonds\n2. Ventes\n3. Clôturer', 'Billetage et total conformes ; « négatif = manquant »', 'Haute')
c(m, 'Gestion de caisse', 'Données du jour à l\'ouverture', '1. Ouvrir Gestion de caisse', 'Données du jour affichées sans recherche', 'Basse')
c(m, 'Gestion de caisse', 'Onglet Analyse des écarts', '1. Ouvrir l\'onglet', 'Écarts par caissier et par jour', 'Basse')
c(m, 'Ticket Z', 'Ticket Z', '1. Éditer le Ticket Z du jour, imprimer', 'Totaux = caisse ; écran sans iframe ; moins de 2 s', 'Haute')
c(m, 'Balance', 'Balance vente / caisse PDF', '1. Choisir la période, imprimer', 'Officine centrée, « Imprimé le … par » en bas à droite, résumé sur deux colonnes', 'Moyenne', ['retours/test-lot-v.js'])
c(m, 'Liste des caisses', 'Impression de la liste', '1. Imprimer un mois', 'PDF en quelques secondes', 'Basse')
c(m, 'Mouvements de caisse', 'Entrée / sortie et édition', '1. Saisir un mouvement\n2. Éditer', 'Mouvement pris dans la clôture ; édition correcte', 'Moyenne')
c(m, 'Journal', 'Journal complet', '1. Filtrer par poste\n2. Afficher les alertes', 'Poste, IP, application ; annulations en série signalées', 'Basse', ['retours/test-journal-complet.js'])

m = module('Facturation et différés', 'FACTURATION')
c(m, 'Différés', 'Régler un différé', '1. Choisir le client\n2. Règlement partiel puis total', 'Solde exact ; contrôles du serveur (appartenance, montant, total)', 'Haute', ['audits/repro-differes.js'])
c(m, 'Différés', 'Modifier une vente différée', '1. Modifier une vente différée non réglée', 'Solde recalculé, aucune dette en double', 'Haute', ['retours/test-differes-modification.js'])
c(m, 'Différés', 'Solde et relevé PDF / Excel', '1. Onglet Solde, période\n2. Imprimer, exporter', 'Montants « 4 000 » avec séparateur ; mêmes chiffres qu\'à l\'écran', 'Moyenne', ['retours/test-differes-solde.js'])
c(m, 'Factures', 'Générer les factures', '1. Générer pour un organisme', 'Numérotation sans doublon, même en concurrence', 'Haute', ['audits/repro-facturation.js'])
c(m, 'Factures', 'Cachet / signature', '1. Imprimer une facture', 'Cachet / signature à la place de « LE PHARMACIEN »', 'Basse')
c(m, 'Factures', 'Icônes de la liste', '1. Ouvrir la liste des factures et des factures de groupe', 'Icônes au trait (supprimer, régler, verrou, certifier)', 'Basse', ['retours/test-style-lots-0510.js'])
c(m, 'Factures subrogatoires', 'Liste et impression', '1. Période, recherche, imprimer', 'Mêmes ventes que l\'ancienne liste ; PDF en onglet', 'Moyenne', ['retours/test-facture-subrogatoire-api.js'])
c(m, 'Règlements', 'Liste des règlements', '1. Filtrer, imprimer', 'Règlements de la période, totaux justes', 'Moyenne', ['reglement/test-liste-reglements.js'])
c(m, 'Récap organisme', 'Récapitulatif par organisme', '1. Choisir l\'organisme et la période', 'Totaux = factures ; édition PDF', 'Basse', ['recap/test-recap-organisme.js'])

m = module('Tiers-payants et clients', 'GESTION DES TIERS-PAYANTS · GESTION CLIENT')
c(m, 'Fiche client', 'Assurances complémentaires', '1. Fiche client > assurances\n2. Modifier taux, matricule, priorité', 'Enregistrés ; priorité unique contrôlée', 'Haute', ['audits/test-ecran-assurances-client-priorites.js'])
c(m, 'Fiche client', 'Nouveau tiers payant principal', '1. Associer un nouveau RO', 'Ancien RO gardé actif, priorités cohérentes', 'Haute', ['audits/test-ecran-nouveau-tp-principal.js', 'audits/repro-assurances-liens.js'])
c(m, 'Vente', 'Modifier le tiers payant d\'une vente', '1. Vente en cours > Modifier le tiers-payant', 'Pas d\'erreur ; part de la vente gardée au même taux', 'Haute', ['audits/test-ecran-modifier-tp-vente-encours.js'])
c(m, 'Vente en dépôt', 'Grille des complémentaires', '1. Vente en dépôt, client\n2. Modifier la grille', 'Mêmes contrôles que la fiche client ; fenêtres qui se ferment', 'Moyenne', ['retours/test-vente-en-depot.js', 'audits/test-ecran-assurances-client-depot.js'])
c(m, 'Clients', 'Recherche « commence par / contient »', '1. Rechercher un client par un morceau du nom', 'Trouvé en « Contient » seulement ; téléphone avec ou sans séparateurs', 'Moyenne', ['retours/test-recherche-client-conseil.js'])
c(m, 'Clients', 'Import et filtre Type client', '1. « Importer »\n2. Filtrer par type', 'Un seul bouton Importer, filtre séparé de la recherche', 'Basse')
c(m, 'Carnet dépôt', 'Gérer carnet dépôt', '1. Choisir un carnet', 'Seuls les carnets en dépôt sont proposés ; recherche lancée', 'Moyenne', ['verif-carnet.js', 'audits/repro-carnet.js'])
c(m, 'Carnet dépôt', 'Retour carnet scanné', '1. Nouveau retour, scanner le produit', 'Code lu encodé (« / » et « # » acceptés)', 'Basse', ['journal/test-retour-carnet-code-encode.js'])
c(m, 'Tiers payants désactivés', 'Liste', '1. Ouvrir, rechercher', 'Liste remplie, recherche OK', 'Basse')
c(m, 'Suivi dossier client', 'Fiche et suivi de consommation', '1. Ouvrir un client\n2. Par vente / par produit', 'Date de naissance et sexe complétables, suivi correct', 'Basse', ['retours/test-ordonnances-fiche-client.js'])

m = module('Commandes et PharmaML', 'GESTION DES COMMANDES')
c(m, 'Suggestion', 'Créer une suggestion (grand volume)', '1. Analyse tiers payant ou articles vendus\n2. « Créer une suggestion »', 'Suggestion créée en moins d\'une minute (4 186 produits : 21 s)', 'Haute', ['analyse/test-analyse-tierspayant.js'])
c(m, 'Suggestion', 'Produit géré en réserve', '1. Vendre un produit en réserve', 'Comparé au seuil avec rayon + réserve ; réassort plafonné au rayon', 'Moyenne', ['retours/test-suggestion-reserve.js'])
c(m, 'Suggestion', 'Équivalents DCI', '1. Ouvrir une suggestion\n2. « Équivalents DCI »', 'Couverture, retrait, reliquat commandé, journal', 'Moyenne', ['retours/test-equivalents-dci-suggestion.js'])
c(m, 'PharmaML', 'Commander par PharmaML', '1. « Commander » au bas de la suggestion', 'Envoi en arrière-plan, statut et réponse sur la ligne', 'Haute', ['retours/test-commander-suggestion-pharmaml.js', 'retours/test-passation-ecrans.js'])
c(m, 'Commandes en cours', 'Onglets Ruptures, Substitutions, Alertes', '1. Ouvrir les onglets', 'Ruptures (heure, motif), substitutions suivies, alertes', 'Moyenne', ['retours/test-commande-en-cours-0825.js'])
c(m, 'Commandes en cours', 'Supprimer une commande', '1. Icône de suppression, 2e clic', 'Confirmation dans l\'icône, pas de fenêtre ; pas d\'erreur au 2e envoi', 'Moyenne')
c(m, 'Commandes en cours', 'Import de commande annulé', '1. Importer, puis Annuler', 'Rien de créé, écran propre', 'Basse', ['support/anomalies-centre/test-import-commande-annuler.js'])
c(m, 'Entrée en stock', 'Entrée d\'un BL', '1. Entrer le BL\n2. Prix, lots, réserve', 'Stocks, prix et lots justes ; suggestion de réserve ouverte au seuil', 'Haute', ['retours/test-entree-stock-0825.js', 'retours/test-reserve-apres-entree.js', 'support/anomalies-centre/test-prix-bl.js'])
c(m, 'Pointage BL', 'Pointer les BL / avoirs', '1. Importer le relevé du grossiste', 'Rapprochement BL / avoirs, colonnes mémorisées par grossiste', 'Moyenne')
c(m, 'Péremptions courtes', 'Écran et exports', '1. Filtrer, changer un statut dans la grille\n2. Excel, PDF', 'Statut enregistré sans fenêtre ; exports corrects', 'Basse', ['retours/test-peremptions-courtes-ecran.js'])
c(m, 'Retours fournisseurs', 'Retour par utilisateur', '1. Filtrer par utilisateur\n2. Tableau de bord, impression', 'Stock jamais négatif ; retours par utilisateur', 'Moyenne', ['points-28-10/test-points-divers.js'])

m = module('Stock', 'GESTION DU STOCK')
c(m, 'Fiche article', 'Détail, courbes, comparaison', '1. Ouvrir un article vendu sur 2 ans\n2. Courbes, « Comparer »', 'Comparaison des années, sections conservées', 'Moyenne', ['ui-fiche-article-refonte.js'])
c(m, 'Fiche article', 'Lots et péremptions', '1. Voir les lots\n2. Supprimer un lot (2 clics)', 'Lot supprimé, liste rechargée ; « Qté restante lot »', 'Moyenne', ['test-lots.js', 'audits/test-fiche-article-qte-restante-lot.js'])
c(m, 'Fiche article', 'Dernière entrée et EAN fabricant', '1. Ouvrir Modifier et Détail', 'Dern. entrée = mise en stock ; EAN en lecture seule au détail', 'Basse', ['retours/test-retours-0810c.js'])
c(m, 'Inventaire', 'Inventaire partiel, saisie, clôture', '1. Créer par lettres du nom\n2. Saisir, clôturer', 'Aucune saisie après clôture ; auteur enregistré', 'Haute')
c(m, 'Ajustement', 'Ajustement du stock', '1. Créer, clôturer', 'Une seule clôture, pas d\'annulation après', 'Moyenne')
c(m, 'Réserve', 'Suggestions de réassort', '1. Ouvrir la réserve\n2. Suggestions', 'Réassort proposé, icônes dessinées', 'Moyenne', ['retours/test-suggestion-reserve.js'])
c(m, 'Emplacements', 'Basculement d\'emplacement', '1. Basculer depuis une ligne', 'Origine en tête, produits actifs seulement', 'Basse', ['retours/test-lot-u.js'])
c(m, 'Périmés', 'Saisie et onglet Analyse', '1. Saisir des périmés\n2. Analyse : objectif, intervalles', 'Taux vs objectif, intervalles = total de la période', 'Moyenne', ['retours/test-perimes-analyse-objectif.js'])
c(m, 'Déconditionnement', 'Historique des mouvements', '1. Vendre au détail une boîte ouverte', 'Chaîne cohérente : ouverture puis vente', 'Moyenne', ['vente/test-historique-deconditionnement.js'])
c(m, 'Comparaison stock', 'Exports et impression', '1. Exporter Excel / CSV, imprimer', 'Quelques secondes (327 s avant)', 'Basse')

m = module('Étiquettes', 'GESTION DU STOCK > Étiquettage Pro')
c(m, 'Étiquettage Pro', 'Imprimer depuis une entrée en stock', '1. Entrée en stock > étiquettes', 'Écran épuré, « Retour à commande en cours »', 'Moyenne', ['retours/test-etiquettes-apres-entree.js'])
c(m, 'Étiquettage Pro', 'Lire une étiquette', '1. « Lire une étiquette », scanner deux fois de suite', 'Aucun chiffre perdu, produit reconnu', 'Haute')
c(m, 'Étiquettage Pro', 'Codes GS1 QR / DataMatrix', '1. Imprimer une planche GS1', 'Codes lisibles, informations à gauche', 'Moyenne', ['retours/test-etiquettes-gs1.js'])
c(m, 'Étiquettage Pro', 'Personnalisation de la disposition', '1. Modifier la disposition, valider\n2. Imprimer', 'Disposition appliquée ; droit réservé à l\'administrateur', 'Moyenne', ['retours/test-etiquettes-personnalisation.js'])
c(m, 'Étiquettage Pro', 'Saisie à la main', '1. Rechercher, choisir le lot, quantité au clavier', 'Enchaînement au clavier sans souris', 'Basse')
c(m, 'Étiquettage Pro', 'Types d\'étiquette', '1. Filtrer, cocher des produits', 'Filtres sur une ligne, mise à jour sélective', 'Basse')
c(m, 'Ancien menu Gestion étiquette', 'Passerelle', '1. Ouvrir l\'ancien menu', '7 derniers jours par défaut ; impression par Étiquettage Pro', 'Basse', ['retours/test-passerelle-etiquettes.js'])

m = module('Analyse et pilotage', 'ANALYSE DE GESTION')
c(m, 'Tableau de bord', 'Tuiles et courbes', '1. Ouvrir le tableau de bord', 'Mêmes chiffres que l\'ancien ; valorisation rayon et total', 'Haute', ['retours/test-tableau-bord.js'])
c(m, 'Tableau de bord', 'Carte d\'affluence', '1. Choisir une semaine', 'Carte jour × heure correcte', 'Basse')
c(m, 'Pilotage', 'Onglets À suivre, Stock, Qualité', '1. Ouvrir chaque onglet', 'Chiffres = base (annulations, péremptions < 6 mois)', 'Moyenne', ['retours/test-pilotage-retours-2009.js', 'retours/test-pilotage-stock-qualite.js'])
c(m, 'Pilotage', 'Rentabilité', '1. Saisir les charges du mois', 'Résultat net calculé', 'Basse')
c(m, 'Audit & rapprochements', 'Analyse, rapports, régularisations', '1. Proposer, commenter, valider', 'Sans commentaire : refus ; appliquées une seule fois, journalisées', 'Haute', ['audit/test-audit-regularisations.js'])
c(m, 'Statistiques', 'Familles, rayons, grossistes', '1. Période sans vente puis normale', 'Aucune erreur ; résultats en quelques secondes', 'Moyenne')
c(m, 'CA par zone / dépôt', 'CA et PDF par laboratoire', '1. Filtres multiples, imprimer', 'Noms des filtres sur l\'édition, montants « 1 234 »', 'Basse', ['retours/test-ca-depot.js'])
c(m, 'Activité des équipes', 'Données et analyses', '1. Filtrer, « Vider les champs »', 'Pastilles stables, infobulles entières', 'Basse')
c(m, 'Anomalies par employé', 'Onglet Anomalies', '1. Ouvrir, filtrer un employé', 'Chaque occurrence et son fil d\'actions', 'Basse', ['productivite/test-anomalies-employes.js'])
c(m, 'Rappels traitement', 'Onglet Analyse et exports', '1. Ouvrir, exporter PDF / CSV', 'Analyse et exports corrects', 'Basse', ['retours/test-rappels-analyse.js'])

m = module('Campagnes et avis clients', 'CAMPAGNES SMS / WHATSAPP · AVIS CLIENTS')
c(m, 'Campagnes', 'Campagne WhatsApp avec pièce jointe', '1. Nouvelle campagne, image ou PDF\n2. Envoyer', 'WhatsApp par défaut, suivi des envois', 'Moyenne')
c(m, 'Campagnes', 'Listes de diffusion et automatiques', '1. Créer une liste\n2. Régler un envoi automatique', '10 onglets ; délai de réponse paramétrable', 'Basse')
c(m, 'Discussions', 'Réponses clients et STOP', '1. Un client répond STOP', 'Fil de discussion, motif STOP visible, plus d\'envoi', 'Haute')
c(m, 'Bannière WhatsApp', 'Session à reconnecter', '1. Couper le service WhatsApp', 'Toast pour l\'administrateur seulement, « Reconnecter » relance', 'Moyenne')
c(m, 'Avis clients', 'Saisie publique et analyse', '1. Saisir un avis par la page publique\n2. Analyse, réglages, affiche', 'Avis enregistré, analyse et affiche PDF', 'Basse')
c(m, 'Sécurité', 'Jetons hors du navigateur', '1. Ouvrir les écrans (outils du navigateur)', 'Aucun jeton WhatsApp ni identifiant SMS côté navigateur', 'Haute')

m = module('Ressources humaines', 'RESSOURCES HUMAINES')
c(m, 'Employés', 'Créer un employé depuis un utilisateur', '1. Nouvel employé, choisir l\'utilisateur', 'Fiche en fenêtre modale', 'Moyenne')
c(m, 'Badges', 'Créer et imprimer un badge', '1. Créer, imprimer (une carte par page)', 'N° PRE-AA-NNNNN-C, QR centré, nom sur une ligne', 'Moyenne')
c(m, 'Badges', 'Régénérer un badge', '1. Régénérer le QR\n2. Scanner l\'ancien', 'Ancien refusé ; réimpression identique sinon', 'Haute')
c(m, 'Pointage', 'Pointer avec l\'application', '1. Pointer par QR + GPS', 'Pointage enregistré avec photo de contrôle', 'Moyenne')
c(m, 'Bornes de pointage', 'Enrôler une borne', '1. Onglet Bornes, enrôler', 'Appels signés acceptés sans session', 'Basse')
c(m, 'Demandes', 'Valider une demande', '1. Profil sans droit puis avec', 'Chacun ne voit que ses demandes ; validation selon le droit', 'Moyenne')
c(m, 'Éditions', 'PDF de chaque onglet', '1. Imprimer chaque onglet', 'PDF lisibles', 'Basse')

m = module('Support et serveur', 'CENTRE DE SUPPORT')
c(m, 'Stabilité du serveur', 'État, arrêts, réglages', '1. Ouvrir l\'écran', 'Tuiles mémoire, base, pool ; courbes ; lexique', 'Moyenne', ['support/test-stabilite-serveur.js'])
c(m, 'Diagnostic & bugs', 'Aucune nouvelle erreur', '1. Après la recette, ouvrir le diagnostic', 'Pas de nouvelle erreur due aux tests', 'Haute', ['retours/test-saisie-ecrans.js'])
c(m, 'Détails de vente', 'Lignes d\'une vente', '1. Ouvrir le détail, actualiser', 'Réponse immédiate, lignes gardées', 'Moyenne', ['support/anomalies-centre/test-details-vente.js'])
c(m, 'Anomalies lot 1 et 4', 'Corrections du centre de support', '1. Rejouer les cas (proforma sans client, findone…)', 'Plus d\'erreur ; réponses lisibles', 'Moyenne', ['support/anomalies-centre/test-lot1-api.js', 'support/anomalies-centre/test-lot1-ecrans.js', 'support/anomalies-centre/test-lot4-ecrans.js'])
c(m, 'Auto-diagnostic', 'Migrations et contrôle de nuit', '1. Ouvrir l\'auto-diagnostic', 'Migration en échec signalée ; pas de fausse alerte quotidienne', 'Basse')
c(m, 'Journaux', 'Aucun secret', '1. Lire server.log après la recette', 'Ni mot de passe, ni jeton, ni texte de message patient', 'Haute')

m = module('Performances et charge', 'Tous menus')
c(m, 'Connexion', '15 postes en même temps', '1. Connecter 15 postes ensemble', 'Tous connectés en quelques secondes (scripts en un seul appel)', 'Haute', ['perf/test-scripts-groupes.js'])
c(m, 'Ventes', 'Ventes simultanées du même article', '1. Deux caisses vendent le même article', 'Stock exact, aucune vente perdue', 'Haute')
c(m, 'Écrans', 'Ouverture des écrans principaux', '1. Ouvrir Ventes, Pilotage, Fiche article', 'Utilisables en moins de 3 s', 'Moyenne', ['retours/test-defilement-ecrans.js'])
c(m, 'Lectures', 'Aucune lecture de plus de 30 s', '1. Parcourir les écrans de liste', 'Toutes les listes en moins de 30 s (lots sans filtre à surveiller)', 'Moyenne', ['retours/test-saisie-api.js'])
c(m, 'Exports', 'Exports lourds', '1. Export ERP d\'une semaine\n2. Comparaison stock', 'Quelques secondes au lieu de plusieurs minutes', 'Basse')
c(m, 'Mobile', 'Application mobile', '1. Infos produit, BL de 500 lignes', 'Réponses en 1 à 3 s', 'Moyenne', ['retours/test-logos-mobile.js', 'retours/test-mobile-54-55.js'])

m = module('Éditions et modèles', 'Toutes éditions')
c(m, 'Modèles d\'état', 'Modèles du dossier des rapports', '1. Comparer le dossier REPORTS du serveur au dépôt', 'balance, analyse d\'inventaire, garde : version du dépôt', 'Haute', ['retours/test-lot-v.js'])
c(m, 'Montants', 'Séparateur des milliers', '1. Imprimer relevé de différés, carnet, récap caisse', '« 4 000 » et non « 4000 »', 'Moyenne', ['retours/test-differes-solde.js', 'retours/test-lot-o.js'])
c(m, 'Dossier PDF', 'Dossier des éditions', '1. Démarrer sur un serveur neuf\n2. Imprimer', 'data/reports/pdf créé automatiquement', 'Moyenne')
c(m, 'Ouverture', 'PDF en onglet', '1. Imprimer depuis différents écrans', 'PDF ouvert dans un onglet, en flux', 'Basse', ['retours/test-lot-j.js'])
c(m, 'Filtres', 'Noms des filtres sur les éditions', '1. Filtrer rayons / familles, imprimer', 'Filtres repris en en-tête', 'Basse', ['retours/test-lot-g.js'])


# ---------------------------------------------------------------- resultats des tests d'ecran (facultatifs)
def lire_resultats(fichier):
    r, courant = {}, None
    if not fichier:
        return r
    for ligne in open(fichier, encoding='utf-8', errors='replace'):
        if ligne.startswith('=== '):
            courant = ligne[4:].strip()
        elif ligne.startswith('--- code ') and courant:
            r[courant] = ligne.split()[2] == '0'
            courant = None
    return r


RES = lire_resultats(RESULTATS)


def test_auto(tests):
    if not tests:
        return 'Manuel', None
    if not RES:
        return 'À jouer', None
    joues = [t for t in tests if t in RES]
    if not joues:
        return 'Non joué', None
    ko = [t for t in joues if not RES[t]]
    return ('KO' if ko else 'OK'), ('KO : ' + ', '.join(re.sub(r'^.*/|\.js$', '', t) for t in ko) if ko else None)


# ---------------------------------------------------------------- classeur
COLS = [('N°', 4), ('Écran', 15), ('Cas à vérifier', 24), ('Étapes', 34), ('Résultat attendu', 36), ('Prio.', 7),
        ('Test auto', 9), ('Statut', 9), ('Testeur · date', 12), ('Observation / ticket', 22)]
NC = len(COLS)
DER = get_column_letter(NC)
ENTETE = 5
bleu = PatternFill('solid', start_color='1F4E78')
clair = PatternFill('solid', start_color='EEF4FA')
jaune = PatternFill('solid', start_color='FFF8DC')
trait = Side(style='thin', color='B4C6E7')
bord = Border(left=trait, right=trait, top=trait, bottom=trait)
VERT = PatternFill('solid', start_color='C6EFCE')
ROUGE = PatternFill('solid', start_color='FFC7CE')
ORANGE = PatternFill('solid', start_color='FFEB9C')


def police(**k):
    return Font(name=POLICE, **k)


def mise_en_page(ws, derniere):
    ws.print_area = 'A1:%s%d' % (DER, derniere)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.page_margins.header = ws.page_margins.footer = 0.12
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = '&A'
    ws.oddFooter.right.text = 'Cahier de recette · ' + BRANCHE
    ws.oddFooter.left.size = ws.oddFooter.right.size = 7
    ws.sheet_view.zoomScale = 100


def feuille_module(wb, nom, menu, lignes):
    ws = wb.create_sheet(nom[:31])
    prem = ENTETE + 1
    der = prem + len(lignes) - 1
    st = '$H$%d:$H$%d' % (prem, der)
    ws['A1'] = 'Cahier de recette · ' + nom
    ws['A1'].font = police(bold=True, size=14, color='1F4E78')
    ws.merge_cells('A1:E1')
    ws['F1'] = 'Menu : ' + menu
    ws['F1'].font = police(size=9, italic=True, color='595959')
    ws.merge_cells('F1:%s1' % DER)
    ws['A2'] = 'Branche : ' + BRANCHE + ' · base de test : copie de travail, jamais la production'
    ws['A2'].font = police(size=9, color='595959')
    ws.merge_cells('A2:%s2' % DER)
    # synthese en une ligne
    # synthese en une ligne : cinq paires libelle / valeur sur les dix colonnes
    synth = [('Cas', '=COUNTA($A$%d:$A$%d)' % (prem, der)), ('OK', '=COUNTIF(%s,"OK")' % st), ('KO', '=COUNTIF(%s,"KO")' % st),
             ('Non testé', '=COUNTA($A$%d:$A$%d)-COUNTA(%s)' % (prem, der, st)),
             ('Auto OK / KO', '=COUNTIF($G$%d:$G$%d,"OK")&" / "&COUNTIF($G$%d:$G$%d,"KO")' % (prem, der, prem, der))]
    for k, (lib, f) in enumerate(synth):
        a = ws.cell(row=3, column=2 * k + 1, value=lib)
        b = ws.cell(row=3, column=2 * k + 2, value=f)
        a.font = police(bold=True, size=9)
        a.alignment = Alignment(horizontal='right')
        b.font = police(size=9, bold=True, color='1F4E78')
        b.alignment = Alignment(horizontal='left')
    ws['A4'] = 'Remplir les cellules jaunes · Statut : OK, KO (décrire, ouvrir un ticket), Bloqué, N/A · Test auto : résultat des tests d\'écran sur la branche'
    ws['A4'].font = police(size=8, italic=True, color='595959')
    ws.merge_cells('A4:%s4' % DER)
    for i, (t, larg) in enumerate(COLS, 1):
        x = ws.cell(row=ENTETE, column=i, value=t)
        x.font = police(bold=True, size=9, color='FFFFFF')
        x.fill = bleu
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        x.border = bord
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[ENTETE].height = 18
    for n, (ecran, cas, etapes, attendu, prio, tests) in enumerate(lignes):
        r = prem + n
        auto, detail = test_auto(tests)
        valeurs = [n + 1, ecran, cas, etapes, attendu, prio, auto, None, None, detail]
        for i, v in enumerate(valeurs, 1):
            x = ws.cell(row=r, column=i, value=v)
            x.font = police(size=8)
            x.alignment = Alignment(vertical='top', wrap_text=True, horizontal='center' if i in (1, 6, 7, 8) else None)
            x.border = bord
            if n % 2:
                x.fill = clair
        ws.cell(row=r, column=3).font = police(size=8, bold=True)
        for i in (8, 9, 10):
            ws.cell(row=r, column=i).fill = jaune
            ws.cell(row=r, column=i).font = police(size=8, color='0000FF')
        if tests:
            ws.cell(row=r, column=7).font = police(size=8, bold=True)
    dv = DataValidation(type='list', formula1='"OK,KO,Bloqué,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('H%d:H%d' % (prem, der))
    for zone in ('H%d:H%d' % (prem, der), 'G%d:G%d' % (prem, der)):
        ws.conditional_formatting.add(zone, CellIsRule(operator='equal', formula=['"OK"'], fill=VERT, font=Font(name=POLICE, color='006100', bold=True)))
        ws.conditional_formatting.add(zone, CellIsRule(operator='equal', formula=['"KO"'], fill=ROUGE, font=Font(name=POLICE, color='9C0006', bold=True)))
        ws.conditional_formatting.add(zone, CellIsRule(operator='equal', formula=['"Bloqué"'], fill=ORANGE, font=Font(name=POLICE, color='9C5700', bold=True)))
    ws.freeze_panes = 'A%d' % prem
    mise_en_page(ws, der)
    return prem, der


par_module = OrderedDict((m, []) for m in MODULES)
for (m, *reste) in CAS:
    par_module[m].append(tuple(reste))

wb = Workbook()
syn = wb.active
syn.title = 'Synthèse'
bornes = []
for m, lignes in par_module.items():
    bornes.append((m,) + feuille_module(wb, m, MODULES[m], lignes))

syn['A1'] = 'Cahier de recette · branche ' + BRANCHE
syn['A1'].font = police(bold=True, size=14, color='1F4E78')
syn.merge_cells('A1:I1')
syn['A2'] = ('Contenu : claude/new-session-xm8ptu + tests-performance + anomalies-support + audits-code (sans l\'inventaire des JSP). '
             'Cas tirés des changements depuis la dernière version de dev (05/10). Tests sur la copie de travail, jamais la base de production.')
syn['A2'].font = police(size=9, color='595959')
syn['A2'].alignment = Alignment(wrap_text=True, vertical='top')
syn.merge_cells('A2:I2')
syn.row_dimensions[2].height = 30
tete = ['Module (feuille)', 'Menu', 'Cas', 'OK', 'KO', 'Bloqué', 'Non testé', 'Tests auto OK / KO', 'Avancement']
largeurs = [28, 40, 7, 7, 7, 8, 10, 16, 11]
for i, (t, l) in enumerate(zip(tete, largeurs), 1):
    x = syn.cell(row=4, column=i, value=t)
    x.font = police(bold=True, size=9, color='FFFFFF')
    x.fill = bleu
    x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    x.border = bord
    syn.column_dimensions[get_column_letter(i)].width = l
for k, (m, prem, der) in enumerate(bornes):
    r = 5 + k
    f = "'%s'!" % m[:31]
    st = '%s$H$%d:$H$%d' % (f, prem, der)
    vals = [m, MODULES[m], '=COUNTA(%s$A$%d:$A$%d)' % (f, prem, der), '=COUNTIF(%s,"OK")' % st, '=COUNTIF(%s,"KO")' % st,
            '=COUNTIF(%s,"Bloqué")' % st, '=C%d-COUNTA(%s)' % (r, st),
            '=COUNTIF(%s$G$%d:$G$%d,"OK")&" / "&COUNTIF(%s$G$%d:$G$%d,"KO")' % (f, prem, der, f, prem, der),
            '=IF(C%d=0,0,(D%d+E%d+F%d)/C%d)' % (r, r, r, r, r)]
    for i, v in enumerate(vals, 1):
        x = syn.cell(row=r, column=i, value=v)
        x.font = police(size=9, bold=(i == 1))
        x.border = bord
        x.alignment = Alignment(horizontal='left' if i <= 2 else 'center', vertical='center')
        if k % 2:
            x.fill = clair
    syn.cell(row=r, column=1).hyperlink = "#'%s'!A1" % m[:31]
    syn.cell(row=r, column=9).number_format = '0%'
tot = 5 + len(bornes)
syn.cell(row=tot, column=1, value='TOTAL').font = police(bold=True, size=9)
for i, col in enumerate('CDEFG', 3):
    x = syn.cell(row=tot, column=i, value='=SUM(%s5:%s%d)' % (col, col, tot - 1))
    x.font = police(bold=True, size=9)
    x.alignment = Alignment(horizontal='center')
x = syn.cell(row=tot, column=9, value='=IF(C%d=0,0,(D%d+E%d+F%d)/C%d)' % ((tot,) * 5))
x.number_format = '0%'
x.font = police(bold=True, size=9)
x.alignment = Alignment(horizontal='center')
for i in range(1, 10):
    syn.cell(row=tot, column=i).border = bord
    syn.cell(row=tot, column=i).fill = PatternFill('solid', start_color='DDEBF7')
syn.print_area = 'A1:I%d' % tot
syn.page_setup.paperSize = syn.PAPERSIZE_A4
syn.page_setup.orientation = 'landscape'
syn.page_setup.fitToWidth = syn.page_setup.fitToHeight = 1
syn.sheet_properties.pageSetUpPr.fitToPage = True
syn.print_options.horizontalCentered = True
syn.oddFooter.right.text = 'Cahier de recette · ' + BRANCHE
syn.oddFooter.right.size = 7
wb.save(SORTIE)
print('%d cas, %d modules -> %s' % (len(CAS), len(MODULES), SORTIE))
