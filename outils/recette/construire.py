# -*- coding: utf-8 -*-
"""Cahier de recette : GENERATEUR du classeur outils/Cahier_de_recette_evolution5.xlsx.

Le classeur n'est pas edite a la main, il est produit par ce script :

    cd outils/recette && python3 construire.py

Dependance : openpyxl (pip install openpyxl). Corriger le classeur directement le
ferait diverger de son contenu source des la vague suivante.

Deux fichiers, deux roles :
  - cas.py         le CONTENU, et rien d'autre : une liste de cas
                   (numero, ecran, action, resultat attendu, reference), et des
                   entrees ('SECTION', titre). C'est ici qu'on ajoute un cas.
  - construire.py  la mise en forme et la mise en page d'impression. Aucun
                   contenu de recette ici.

CE QUE LA MISE EN PAGE GARANTIT, ET POURQUOI

  - UNE SEULE feuille : demande de l'officine.
  - A4 paysage, ajuste a UNE page de large (fitToWidth=1, fitToHeight=0) : le
    tableau ne se coupe jamais en deux dans la largeur, quel que soit le nombre
    de cas.
  - En-tetes de colonnes REPETES sur chaque page (print_title_rows). Sans cela,
    la page 4 n'est qu'une grille de cases sans titre.
  - Tout tient dans les colonnes A a G, c'est-a-dire dans la zone imprimee. Une
    premiere version posait deux valeurs du recapitulatif en colonne H : elles
    n'apparaissaient pas sur le papier.
  - Hauteurs de ligne GENEREUSES (50 et 56 caracteres par ligne estimes, la ou
    la largeur en laisserait passer davantage) : Excel ROGNE le texte renvoye a
    la ligne des que la hauteur imposee est trop juste, et une consigne de
    recette coupee en plein milieu ne se devine pas.
  - Recapitulatif par FORMULES (COUNTIF sur les verdicts, COUNTA sur les
    references) : aucun nombre ecrit en dur, le compte suit ce que le recetteur
    saisit. Liste de choix OK / KO / N/A sur la colonne Verdict.

UNE LIMITE CONNUE

openpyxl ecrit les formules sans valeur calculee : tant qu'Excel n'a pas ouvert le
fichier, les cinq cellules du recapitulatif se lisent vides par un programme.
Excel recalcule a l'ouverture, il n'y a donc rien a faire pour le recetteur.

Le recalcul de controle par LibreOffice n'a pas pu etre passe sur le conteneur qui
a produit ce classeur : LibreOffice y refuse de charger le moindre fichier, y
compris un CSV trivial. Les plages des formules ont donc ete verifiees autrement,
en evaluant les memes comptages en Python sur le classeur produit. Sur un poste ou
LibreOffice fonctionne, lancer le recalcul avant diffusion.
"""
import importlib.util
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

ICI = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('cas', os.path.join(ICI, 'cas.py'))
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)
CAS = mod.CAS

POLICE = 'Arial'
BRANCHE = 'claude/new-session-xm8ptu'
COMMITS = '26 commits (c622289 -> 9d214eb)'

# Couleurs : sobres, lisibles en noir et blanc comme en couleur.
BLEU = '1F4E79'
BLEU_PALE = 'DCE6F1'
GRIS_PALE = 'F2F2F2'
VERT_PALE = 'E2EFDA'

fin = Side(style='thin', color='BFBFBF')
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)

wb = Workbook()
ws = wb.active
ws.title = 'Cahier de recette'

# ---------------------------------------------------------------- colonnes
# Largeurs choisies pour tenir sur UNE page de large en A4 paysage.
LARGEURS = {'A': 6.5, 'B': 20, 'C': 47, 'D': 53, 'E': 11, 'F': 8.5, 'G': 22}
for col, largeur in LARGEURS.items():
    ws.column_dimensions[col].width = largeur

HAUT = Alignment(vertical='top', wrap_text=True)
HAUT_CENTRE = Alignment(vertical='top', horizontal='center', wrap_text=True)

ligne = 1

# ---------------------------------------------------------------- titre
ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=7)
c = ws.cell(row=ligne, column=1, value='CAHIER DE RECETTE — PRESTIGE — ÉVOLUTION 5 ET RETOURS DU 17/09')
c.font = Font(name=POLICE, size=14, bold=True, color='FFFFFF')
c.fill = PatternFill('solid', fgColor=BLEU)
c.alignment = Alignment(vertical='center', horizontal='center')
ws.row_dimensions[ligne].height = 26
ligne += 1

ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=7)
c = ws.cell(row=ligne, column=1,
            value='Branche %s — %s.   Recette effectuée par : ..............................'
                  '     Date : ......... / ......... / .........     Version déployée : ..............................'
                  % (BRANCHE, COMMITS))
c.font = Font(name=POLICE, size=9, italic=True)
c.alignment = Alignment(vertical='center', horizontal='center')
ws.row_dimensions[ligne].height = 18
ligne += 1

# ---------------------------------------------------------------- mode d'emploi + recap
# Bandeau d'en-tete : mode d'emploi a gauche (A:D), recapitulatif a droite (E:F).
# TOUT tient dans les colonnes A a G, c'est-a-dire dans la zone imprimee : une
# premiere version posait deux valeurs en colonne H, qui ne sortait pas a l'impression.
HAUTEUR_BANDEAU = 5
ligne_recap = ligne
ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne + HAUTEUR_BANDEAU - 1, end_column=4)
c = ws.cell(row=ligne, column=1,
            value="MODE D'EMPLOI\n"
                  "Renseigner la colonne « Verdict » pour chaque cas : OK, KO, ou N/A lorsque le cas ne s'applique pas "
                  "à cette officine (une liste de choix est proposée dans la cellule).\n"
                  "Toute anomalie se note dans « Observations » : ce qui a été fait, et ce qui s'est passé.\n"
                  "La section A pose les prérequis : sans eux, les cas des sections C et D ne sont pas vérifiables.\n"
                  "La colonne « Réf. » donne le commit d'origine, pour retrouver le détail de la modification.")
c.font = Font(name=POLICE, size=9)
c.alignment = Alignment(vertical='top', wrap_text=True)
for r in range(ligne, ligne + HAUTEUR_BANDEAU):
    for col in range(1, 5):
        ws.cell(row=r, column=col).border = BORD
        ws.cell(row=r, column=col).fill = PatternFill('solid', fgColor=GRIS_PALE)
    ws.row_dimensions[r].height = 13

# Le recapitulatif est renseigne par des FORMULES : il suit les verdicts saisis.
RECAP = ['Cas au total', 'OK', 'KO', 'N/A', 'Reste à faire']
cellules_recap = {}
for i, libelle in enumerate(RECAP):
    r = ligne + i
    cl = ws.cell(row=r, column=5, value=libelle)
    cl.font = Font(name=POLICE, size=9, bold=True)
    cl.alignment = Alignment(vertical='center', wrap_text=True)
    cl.fill = PatternFill('solid', fgColor=BLEU_PALE)
    cl.border = BORD
    cv = ws.cell(row=r, column=6)
    cv.font = Font(name=POLICE, size=10, bold=True)
    cv.alignment = Alignment(vertical='center', horizontal='center')
    cv.fill = PatternFill('solid', fgColor=BLEU_PALE)
    cv.border = BORD
    cellules_recap[libelle] = 'F%d' % r

# La colonne Observations sert ici de zone de synthese, sur toute la hauteur du bandeau.
ws.merge_cells(start_row=ligne, start_column=7, end_row=ligne + HAUTEUR_BANDEAU - 1, end_column=7)
c = ws.cell(row=ligne, column=7, value='Synthèse / réserves :')
c.font = Font(name=POLICE, size=9, italic=True)
c.alignment = Alignment(vertical='top', wrap_text=True)
c.border = BORD
for r in range(ligne, ligne + HAUTEUR_BANDEAU):
    ws.cell(row=r, column=7).border = BORD

ligne += HAUTEUR_BANDEAU

ligne += 1  # une ligne de respiration

# ---------------------------------------------------------------- en-tetes
ligne_entetes = ligne
ENTETES = ['N°', 'Écran / Fonction', 'Ce qu\'il faut faire', 'Résultat attendu', 'Réf.', 'Verdict', 'Observations']
for i, titre in enumerate(ENTETES, start=1):
    c = ws.cell(row=ligne, column=i, value=titre)
    c.font = Font(name=POLICE, size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=BLEU)
    c.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
    c.border = BORD
ws.row_dimensions[ligne].height = 22
ligne += 1

premiere_ligne_cas = ligne
lignes_cas = []

for entree in CAS:
    if entree[0] == 'SECTION':
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=7)
        c = ws.cell(row=ligne, column=1, value=entree[1])
        c.font = Font(name=POLICE, size=10, bold=True, color=BLEU)
        c.fill = PatternFill('solid', fgColor=BLEU_PALE)
        c.alignment = Alignment(vertical='center')
        for col in range(1, 8):
            ws.cell(row=ligne, column=col).border = BORD
        ws.row_dimensions[ligne].height = 19
        ligne += 1
        continue

    numero, ecran, action, attendu, ref = entree
    ws.cell(row=ligne, column=1, value=numero).font = Font(name=POLICE, size=9, bold=True)
    ws.cell(row=ligne, column=2, value=ecran).font = Font(name=POLICE, size=9, bold=True)
    ws.cell(row=ligne, column=3, value=action).font = Font(name=POLICE, size=9)
    ws.cell(row=ligne, column=4, value=attendu).font = Font(name=POLICE, size=9)
    ws.cell(row=ligne, column=5, value=ref).font = Font(name=POLICE, size=8, color='808080')
    ws.cell(row=ligne, column=6, value=None)
    ws.cell(row=ligne, column=7, value=None)
    for col in range(1, 8):
        cel = ws.cell(row=ligne, column=col)
        cel.alignment = HAUT_CENTRE if col in (1, 5, 6) else HAUT
        cel.border = BORD
        if col == 6:
            cel.fill = PatternFill('solid', fgColor='FFF2CC')  # la colonne a remplir
    # Hauteur de ligne : estimee sur le texte le plus long des deux colonnes larges.
    #
    # Volontairement GENEREUSE (50 et 56 caracteres par ligne, la ou la largeur en
    # laisserait passer davantage) : Excel ROGNE le texte renvoye a la ligne des que la
    # hauteur imposee est trop juste, et une consigne de recette coupee en plein milieu
    # ne se devine pas. Mieux vaut un blanc en bas de cellule qu'une phrase tronquee.
    lignes_action = max(1, -(-len(action) // 50))
    lignes_attendu = max(1, -(-len(attendu) // 56))
    ws.row_dimensions[ligne].height = max(26, 12.0 * max(lignes_action, lignes_attendu) + 6)
    lignes_cas.append(ligne)
    ligne += 1

derniere_ligne_cas = ligne - 1

# ---------------------------------------------------------------- formules du recap
plage = 'F%d:F%d' % (premiere_ligne_cas, derniere_ligne_cas)
# Le total est COMPTE par formule et non ecrit en dur : la colonne « Ref. » est
# renseignee sur chaque cas et vide sur les lignes de section, fusionnees de A a G.
ws[cellules_recap['Cas au total']] = '=COUNTA(E%d:E%d)' % (premiere_ligne_cas, derniere_ligne_cas)
ws[cellules_recap['OK']] = '=COUNTIF(%s,"OK")' % plage
ws[cellules_recap['KO']] = '=COUNTIF(%s,"KO")' % plage
ws[cellules_recap['N/A']] = '=COUNTIF(%s,"N/A")' % plage
ws[cellules_recap['Reste à faire']] = '=%s-%s-%s-%s' % (
    cellules_recap['Cas au total'], cellules_recap['OK'],
    cellules_recap['KO'], cellules_recap['N/A'])

# ---------------------------------------------------------------- liste de choix du verdict
validation = DataValidation(type='list', formula1='"OK,KO,N/A"', allow_blank=True, showDropDown=False)
validation.error = 'Verdict attendu : OK, KO ou N/A.'
validation.errorTitle = 'Verdict'
validation.prompt = 'OK, KO ou N/A'
ws.add_data_validation(validation)
validation.add('F%d:F%d' % (premiere_ligne_cas, derniere_ligne_cas))

# ---------------------------------------------------------------- pied de page du document
ligne += 1
ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=7)
c = ws.cell(row=ligne, column=1,
            value="Vérifications automatisées existantes pour cette branche — suite unitaire : 901 cas. "
                  "Bout en bout (src/test/e2e/retours) : test-depot-onglets 57, test-depot-extension-stock 55, "
                  "test-vente-en-depot 31, test-posos 32, test-client-standard 32, test-import-clients 32, "
                  "test-stock-reserve-colonnes 39, test-vente-contexte-depot 26, test-ca-depot 20, "
                  "test-eclater-suggestion 19, test-redimensionnement 14, et les suites des points 2, 5, 6 et 8. "
                  "Ces tests ne remplacent pas la recette : ils protègent des régressions, la recette valide l'usage.")
c.font = Font(name=POLICE, size=8, italic=True)
c.alignment = Alignment(vertical='top', wrap_text=True)
c.fill = PatternFill('solid', fgColor=VERT_PALE)
for col in range(1, 8):
    ws.cell(row=ligne, column=col).border = BORD
    ws.cell(row=ligne, column=col).fill = PatternFill('solid', fgColor=VERT_PALE)
ws.row_dimensions[ligne].height = 40
ligne += 2

ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=7)
c = ws.cell(row=ligne, column=1,
            value='Recette prononcée le ......... / ......... / .........          '
                  'Visa du pharmacien : ..............................................          '
                  'Réserves : ..............................................................................')
c.font = Font(name=POLICE, size=9, bold=True)
c.alignment = Alignment(vertical='center')
ws.row_dimensions[ligne].height = 30

# ---------------------------------------------------------------- mise en page pour l'impression
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0          # autant de pages en hauteur qu'il faut
ws.page_margins.left = 0.35
ws.page_margins.right = 0.3
ws.page_margins.top = 0.45
ws.page_margins.bottom = 0.5
ws.page_margins.header = 0.2
ws.page_margins.footer = 0.2
ws.print_options.horizontalCentered = True
# Les en-tetes de colonnes se repetent sur chaque page : sans cela, la page 4
# est une grille de cases sans titre.
ws.print_title_rows = '%d:%d' % (ligne_entetes, ligne_entetes)
ws.print_area = 'A1:G%d' % ligne
ws.oddFooter.left.text = 'Cahier de recette — %s' % BRANCHE
ws.oddFooter.left.size = 8
ws.oddFooter.left.font = POLICE
ws.oddFooter.right.text = 'Page &P / &N'
ws.oddFooter.right.size = 8
ws.oddFooter.right.font = POLICE
# A l'ecran : les en-tetes restent visibles au defilement.
ws.freeze_panes = 'A%d' % (ligne_entetes + 1)
ws.sheet_view.zoomScale = 100

# Le classeur se depose toujours dans outils/, quel que soit le repertoire d'appel.
SORTIE = os.path.join(os.path.dirname(ICI), 'Cahier_de_recette_evolution5.xlsx')
wb.save(SORTIE)
print('classeur :', SORTIE)
print('cas : %d, lignes %d a %d' % (len(lignes_cas), premiere_ligne_cas, derniere_ligne_cas))
print('en-tetes repetes : ligne %d' % ligne_entetes)
